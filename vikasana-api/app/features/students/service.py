# app/features/students/service.py

import os
import csv
import io
from typing import List, Tuple

from openpyxl import load_workbook

from sqlalchemy import select, or_
from sqlalchemy.ext.asyncio import AsyncSession

import asyncio
from app.features.faculty.models import Faculty
from app.features.faculty.role_policy import (
    ROLE_FACULTY,
    ROLE_HOD,
    normalize_faculty_role,
)
from app.features.students.models import Student, StudentType
from app.features.students.schemas.student import StudentCreate
from app.core.email_service import send_student_welcome_email


def _clean(v: str) -> str:
    return (v or "").strip()


def _parse_student_type(v: str) -> StudentType:
    x = (v or "").strip().upper()
    if x in ("DIPLOMA", "DIPLOMA_SCHEME", "DIPLOMA SCHEME"):
        return StudentType.DIPLOMA
    return StudentType.REGULAR


def _required_points_for_type(stype: StudentType) -> int:
    return 60 if stype == StudentType.DIPLOMA else 100


def _is_active_student(s: Student) -> bool:
    return bool(getattr(s, "is_active", True))


def _reactivate_student(
    s: Student,
    *,
    college: str,
    name: str,
    email: str,
    usn: str,
    branch: str,
    student_type: StudentType,
    passout_year: int,
    admitted_year: int,
    faculty_id: int | None,
) -> Student:
    """
    Permanent fix:
    If admin previously deleted/deactivated a student, CSV/manual add should restore
    that same row instead of creating duplicate or skipping it.
    """
    s.is_active = True
    s.college = college
    s.name = name
    s.email = email
    s.usn = usn
    s.branch = branch
    s.student_type = student_type
    s.required_total_points = _required_points_for_type(student_type)
    s.passout_year = passout_year
    s.admitted_year = admitted_year

    # Preserve original creator provenance on reactivation.
    # Only backfill legacy rows where creator provenance is missing.
    if s.created_by_faculty_id is None:
        s.created_by_faculty_id = faculty_id

    # Do NOT reset earned points/certificates here.
    # Old activity/certificate history remains attached to the same student id.
    return s


def _trigger_student_welcome_email(email: str | None, name: str | None):
    if not email:
        return

    async def _send_later():
        try:
            app_url = os.getenv(
                "STUDENT_APP_DOWNLOAD_URL",
                "https://loraa-connect.vikasanafoundation.org/install",
            )
            await send_student_welcome_email(
                to_email=email,
                to_name=name or "Student",
                app_download_url=app_url,
            )
        except Exception as e:
            print(f"[WARN] Student welcome email not sent for {email}: {e}")

    try:
        asyncio.create_task(_send_later())
    except RuntimeError:
        print(f"[WARN] Could not schedule welcome email for {email}")


def _coerce_student_type(v) -> StudentType:
    if isinstance(v, StudentType):
        return v
    return _parse_student_type(str(v))


def _normalize_csv_headers(fieldnames: list[str] | None) -> tuple[dict[str, str], set[str]]:
    if not fieldnames:
        return {}, set()

    field_map = {h.strip().lower(): h for h in fieldnames if h and h.strip()}
    return field_map, set(field_map.keys())



def record_student_assignment_history(
    *,
    db: AsyncSession,
    student: Student,
    previous_department_id: int | None,
    new_department_id: int | None,
    previous_batch_id: int | None,
    new_batch_id: int | None,
    previous_year: int | None,
    new_year: int | None,
    previous_faculty_id: int | None,
    new_faculty_id: int | None,
    changed_by_faculty_id: int | None = None,
    changed_by_admin_id: int | None = None,
    reason: str | None = None,
):
    """
    Add academic and faculty-assignment history records when values
    genuinely change.

    This helper does not flush or commit the database transaction.
    """
    from app.features.organization.models import (
        StudentAcademicHistory,
        StudentFacultyAssignment,
    )

    academic_history = None
    faculty_history = None

    academic_changed = any(
        previous != current
        for previous, current in (
            (
                previous_department_id,
                new_department_id,
            ),
            (
                previous_batch_id,
                new_batch_id,
            ),
            (
                previous_year,
                new_year,
            ),
        )
    )

    if academic_changed:
        previous_values = (
            previous_department_id,
            previous_batch_id,
            previous_year,
        )
        new_values = (
            new_department_id,
            new_batch_id,
            new_year,
        )

        if (
            all(value is None for value in previous_values)
            and any(value is not None for value in new_values)
        ):
            academic_action = "ACADEMIC_ASSIGNED"
        elif (
            any(value is not None for value in previous_values)
            and all(value is None for value in new_values)
        ):
            academic_action = "ACADEMIC_CLEARED"
        else:
            academic_action = "ACADEMIC_UPDATED"

        academic_history = StudentAcademicHistory(
            student=student,
            action=academic_action,
            from_department_id=previous_department_id,
            to_department_id=new_department_id,
            from_batch_id=previous_batch_id,
            to_batch_id=new_batch_id,
            from_year=previous_year,
            to_year=new_year,
            reason=reason,
            changed_by_faculty_id=changed_by_faculty_id,
            changed_by_admin_id=changed_by_admin_id,
        )

        db.add(academic_history)

    if previous_faculty_id != new_faculty_id:
        if (
            previous_faculty_id is None
            and new_faculty_id is not None
        ):
            faculty_action = "ASSIGNED"
        elif (
            previous_faculty_id is not None
            and new_faculty_id is None
        ):
            faculty_action = "UNASSIGNED"
        else:
            faculty_action = "REASSIGNED"

        faculty_history = StudentFacultyAssignment(
            student=student,
            previous_faculty_id=previous_faculty_id,
            assigned_faculty_id=new_faculty_id,
            action=faculty_action,
            assigned_by_faculty_id=changed_by_faculty_id,
            assigned_by_admin_id=changed_by_admin_id,
            note=reason,
        )

        db.add(faculty_history)

    return academic_history, faculty_history


def _parse_optional_csv_int(
    *,
    row: dict,
    field_map: dict[str, str],
    headers: set[str],
    field_name: str,
    minimum: int = 1,
    maximum: int | None = None,
) -> int | None:
    """Parse an optional integer CSV field with clear row errors."""
    if field_name not in headers:
        return None

    raw_value = _clean(
        row.get(field_map[field_name], "")
    )

    if not raw_value:
        return None

    try:
        value = int(raw_value)
    except (TypeError, ValueError) as exc:
        raise ValueError(
            f"{field_name} must be an integer"
        ) from exc

    if value < minimum:
        raise ValueError(
            f"{field_name} must be at least {minimum}"
        )

    if maximum is not None and value > maximum:
        raise ValueError(
            f"{field_name} must not exceed {maximum}"
        )

    return value


async def resolve_student_academic_batch_from_years(
    *,
    db: AsyncSession,
    college: str,
    admitted_year: int,
    passout_year: int,
    requested_batch_id: int | None = None,
    requested_current_year: int | None = None,
):
    """
    Resolve the student's active college-wide AcademicBatch from:

        authenticated/student college
        + admitted_year
        + passout_year

    batch_id is therefore server-derived.

    New students without an explicit current_year start in Year 1.
    Existing/reactivated students may preserve their current year
    later in create_student() when they remain in the same batch.

    This helper performs read queries only.
    """

    from sqlalchemy import func, select

    from app.features.organization.models import (
        AcademicBatch,
    )

    college_value = _clean(
        college
    )

    if not college_value:
        raise ValueError(
            "Student college is required for batch assignment"
        )

    admitted_year = int(
        admitted_year
    )

    passout_year = int(
        passout_year
    )

    if passout_year <= admitted_year:
        raise ValueError(
            "Passout year must be after admission year"
        )

    college_key = (
        college_value.casefold()
    )

    result = await db.execute(
        select(
            AcademicBatch
        )
        .where(
            func.lower(
                func.trim(
                    AcademicBatch.college
                )
            )
            == college_key,

            AcademicBatch.admitted_year
            == admitted_year,

            AcademicBatch.passout_year
            == passout_year,

            AcademicBatch.is_active.is_(
                True
            ),
        )
        .order_by(
            AcademicBatch.id.asc()
        )
        .limit(2)
    )

    matches = list(
        result.scalars().all()
    )

    batch_label = (
        f"{admitted_year}"
        f"–"
        f"{passout_year}"
    )

    if not matches:
        raise ValueError(
            f"Academic batch {batch_label} "
            f"has not been created for {college_value}. "
            "Create the batch in Batch Management first."
        )

    if len(matches) > 1:
        raise ValueError(
            f"Multiple active academic batches match "
            f"{batch_label} for {college_value}"
        )

    batch = matches[0]

    # The browser/CSV must never override the server-derived
    # college batch with another batch ID.
    if (
        requested_batch_id is not None
        and int(
            requested_batch_id
        )
        != int(
            batch.id
        )
    ):
        raise ValueError(
            "Provided batch_id does not match the "
            "student admission/passout years"
        )

    duration = int(
        batch.course_duration_years
        or (
            passout_year
            - admitted_year
        )
    )

    if requested_current_year is None:
        resolved_current_year = 1

    else:
        resolved_current_year = int(
            requested_current_year
        )

        if (
            resolved_current_year < 1
            or resolved_current_year > duration
        ):
            raise ValueError(
                "current_year must be between 1 and "
                f"{duration} for academic batch "
                f"{batch_label}"
            )

    return (
        batch,
        resolved_current_year,
    )


async def validate_student_academic_assignment(
    *,
    db: AsyncSession,
    college: str,
    department_id: int | None = None,
    batch_id: int | None = None,
    current_year: int | None = None,
    assigned_faculty_id: int | None = None,
    require_hierarchy_faculty: bool = False,
):
    """
    Validate optional department-wise student assignments.

    This helper performs only read queries. It does not modify or commit
    student, department, batch, or faculty records.
    """
    from fastapi import HTTPException
    from sqlalchemy import select

    from app.features.faculty.models import Faculty
    from app.features.organization.models import (
        AcademicBatch,
        Department,
    )

    college_value = _clean(college)

    if not college_value:
        raise HTTPException(
            status_code=400,
            detail="Student college is required for academic assignment",
        )

    if current_year is not None and not 1 <= current_year <= 8:
        raise HTTPException(
            status_code=400,
            detail="current_year must be between 1 and 8",
        )

    department = None
    batch = None
    assigned_faculty = None

    if department_id is not None:
        department_result = await db.execute(
            select(Department).where(
                Department.id == department_id,
                Department.college == college_value,
                Department.is_active.is_(True),
            )
        )
        department = department_result.scalar_one_or_none()

        if department is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Department must be active and belong to the "
                    "same college as the student"
                ),
            )

    if batch_id is not None:
        batch_result = await db.execute(
            select(AcademicBatch).where(
                AcademicBatch.id == batch_id,
                AcademicBatch.college == college_value,
                AcademicBatch.is_active.is_(True),
            )
        )
        batch = batch_result.scalar_one_or_none()

        if batch is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Academic batch must be active and belong to the "
                    "same college as the student"
                ),
            )

        duration = getattr(
            batch,
            "course_duration_years",
            None,
        )

        if (
            current_year is not None
            and duration is not None
            and current_year > duration
        ):
            raise HTTPException(
                status_code=400,
                detail=(
                    "current_year cannot exceed the selected batch "
                    "course duration"
                ),
            )

    if assigned_faculty_id is not None:
        if department_id is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "department_id is required when "
                    "assigned_faculty_id is provided"
                ),
            )

        faculty_result = await db.execute(
            select(Faculty).where(
                Faculty.id == assigned_faculty_id,
                Faculty.college == college_value,
                Faculty.is_active.is_(True),
            )
        )
        assigned_faculty = (
            faculty_result.scalar_one_or_none()
        )

        if assigned_faculty is None:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Assigned faculty must be active and belong to "
                    "the same college as the student"
                ),
            )

        if assigned_faculty.department_id != department_id:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Assigned faculty must belong to the student's "
                    "selected department"
                ),
            )

        if require_hierarchy_faculty:
            try:
                assigned_role = normalize_faculty_role(
                    assigned_faculty.role
                )
            except ValueError as error:
                raise HTTPException(
                    status_code=400,
                    detail="Assigned account has an invalid Faculty role",
                ) from error

            if assigned_role != ROLE_FACULTY:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "assigned_faculty_id must reference a "
                        "Faculty/Mentor account"
                    ),
                )

            parent_faculty_id = getattr(
                assigned_faculty,
                "parent_faculty_id",
                None,
            )

            if parent_faculty_id is None:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Assigned Faculty/Mentor must be directly "
                        "under an active HOD"
                    ),
                )

            parent_result = await db.execute(
                select(Faculty).where(
                    Faculty.id == parent_faculty_id,
                    Faculty.college == college_value,
                    Faculty.department_id == department_id,
                    Faculty.is_active.is_(True),
                )
            )

            parent_faculty = parent_result.scalar_one_or_none()

            if parent_faculty is None:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Assigned Faculty/Mentor must be directly "
                        "under an active HOD in the same college "
                        "and department"
                    ),
                )

            try:
                parent_role = normalize_faculty_role(
                    parent_faculty.role
                )
            except ValueError as error:
                raise HTTPException(
                    status_code=400,
                    detail="Assigned Faculty/Mentor parent has an invalid role",
                ) from error

            if parent_role != ROLE_HOD:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        "Assigned Faculty/Mentor parent must be an HOD"
                    ),
                )

    return department, batch, assigned_faculty



async def apply_student_assignment_update(
    *,
    db: AsyncSession,
    student: Student,
    assignment_data: dict[str, int | None],
    changed_by_faculty_id: int | None = None,
    changed_by_admin_id: int | None = None,
    reason: str = "Student assignment workflow update",
) -> Student:
    """
    Apply a validated academic/faculty assignment update.

    The function mutates the supplied Student ORM object and adds
    assignment-history records through the existing Stage 5 helper.

    It deliberately does not commit or refresh the session so that
    single and bulk assignment routes can control the transaction.
    """
    from fastapi import HTTPException

    academic_fields = {
        "department_id",
        "batch_id",
        "current_year",
        "assigned_faculty_id",
    }

    if not assignment_data:
        raise HTTPException(
            status_code=400,
            detail="No student assignment fields provided",
        )

    unsupported_fields = (
        set(assignment_data) - academic_fields
    )

    if unsupported_fields:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported student assignment fields: "
                + ", ".join(sorted(unsupported_fields))
            ),
        )

    previous_department_id = student.department_id
    previous_batch_id = student.batch_id
    previous_current_year = student.current_year
    previous_assigned_faculty_id = (
        student.assigned_faculty_id
    )

    target_department_id = assignment_data.get(
        "department_id",
        student.department_id,
    )
    target_batch_id = assignment_data.get(
        "batch_id",
        student.batch_id,
    )
    target_current_year = assignment_data.get(
        "current_year",
        student.current_year,
    )
    target_assigned_faculty_id = assignment_data.get(
        "assigned_faculty_id",
        student.assigned_faculty_id,
    )

    has_academic_assignment = any(
        value is not None
        for value in (
            target_department_id,
            target_batch_id,
            target_current_year,
            target_assigned_faculty_id,
        )
    )

    if has_academic_assignment:
        await validate_student_academic_assignment(
            db=db,
            college=student.college,
            department_id=target_department_id,
            batch_id=target_batch_id,
            current_year=target_current_year,
            assigned_faculty_id=target_assigned_faculty_id,
        )

    for field_name, value in assignment_data.items():
        setattr(student, field_name, value)

    normalized_reason = str(reason or "").strip()

    if not normalized_reason:
        normalized_reason = (
            "Student assignment workflow update"
        )

    record_student_assignment_history(
        db=db,
        student=student,
        previous_department_id=previous_department_id,
        new_department_id=student.department_id,
        previous_batch_id=previous_batch_id,
        new_batch_id=student.batch_id,
        previous_year=previous_current_year,
        new_year=student.current_year,
        previous_faculty_id=previous_assigned_faculty_id,
        new_faculty_id=student.assigned_faculty_id,
        changed_by_faculty_id=changed_by_faculty_id,
        changed_by_admin_id=changed_by_admin_id,
        reason=normalized_reason,
    )

    return student


async def apply_bulk_student_assignment_update(
    *,
    db: AsyncSession,
    students: list[Student],
    assignment_data: dict[str, int | None],
    changed_by_faculty_id: int | None = None,
    changed_by_admin_id: int | None = None,
    reason: str = "Bulk student assignment workflow update",
) -> list[Student]:
    """
    Apply the same assignment update to multiple students.

    All target assignments are validated before any Student object
    is mutated. This prevents partial in-memory updates when a later
    student fails validation.

    The function adds assignment-history records but deliberately
    does not commit or refresh the database session.
    """
    from fastapi import HTTPException

    academic_fields = {
        "department_id",
        "batch_id",
        "current_year",
        "assigned_faculty_id",
    }

    if not students:
        raise HTTPException(
            status_code=400,
            detail="No students provided for bulk assignment",
        )

    if not assignment_data:
        raise HTTPException(
            status_code=400,
            detail="No student assignment fields provided",
        )

    unsupported_fields = (
        set(assignment_data) - academic_fields
    )

    if unsupported_fields:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported student assignment fields: "
                + ", ".join(sorted(unsupported_fields))
            ),
        )

    student_ids = [
        getattr(student, "id", None)
        for student in students
    ]

    concrete_student_ids = [
        student_id
        for student_id in student_ids
        if student_id is not None
    ]

    if len(concrete_student_ids) != len(
        set(concrete_student_ids)
    ):
        raise HTTPException(
            status_code=400,
            detail=(
                "Duplicate students provided for "
                "bulk assignment"
            ),
        )

    assignment_targets = []

    # First pass: calculate and validate every final state.
    # No Student object is mutated during this pass.
    for student in students:
        target_department_id = assignment_data.get(
            "department_id",
            student.department_id,
        )
        target_batch_id = assignment_data.get(
            "batch_id",
            student.batch_id,
        )
        target_current_year = assignment_data.get(
            "current_year",
            student.current_year,
        )
        target_assigned_faculty_id = assignment_data.get(
            "assigned_faculty_id",
            student.assigned_faculty_id,
        )

        has_academic_assignment = any(
            value is not None
            for value in (
                target_department_id,
                target_batch_id,
                target_current_year,
                target_assigned_faculty_id,
            )
        )

        if has_academic_assignment:
            await validate_student_academic_assignment(
                db=db,
                college=student.college,
                department_id=target_department_id,
                batch_id=target_batch_id,
                current_year=target_current_year,
                assigned_faculty_id=(
                    target_assigned_faculty_id
                ),
            )

        assignment_targets.append(
            (
                student,
                target_department_id,
                target_batch_id,
                target_current_year,
                target_assigned_faculty_id,
            )
        )

    normalized_reason = str(reason or "").strip()

    if not normalized_reason:
        normalized_reason = (
            "Bulk student assignment workflow update"
        )

    # Second pass: mutate only after every student passed validation.
    for (
        student,
        target_department_id,
        target_batch_id,
        target_current_year,
        target_assigned_faculty_id,
    ) in assignment_targets:
        previous_department_id = student.department_id
        previous_batch_id = student.batch_id
        previous_current_year = student.current_year
        previous_assigned_faculty_id = (
            student.assigned_faculty_id
        )

        student.department_id = target_department_id
        student.batch_id = target_batch_id
        student.current_year = target_current_year
        student.assigned_faculty_id = (
            target_assigned_faculty_id
        )

        record_student_assignment_history(
            db=db,
            student=student,
            previous_department_id=(
                previous_department_id
            ),
            new_department_id=student.department_id,
            previous_batch_id=previous_batch_id,
            new_batch_id=student.batch_id,
            previous_year=previous_current_year,
            new_year=student.current_year,
            previous_faculty_id=(
                previous_assigned_faculty_id
            ),
            new_faculty_id=student.assigned_faculty_id,
            changed_by_faculty_id=(
                changed_by_faculty_id
            ),
            changed_by_admin_id=changed_by_admin_id,
            reason=normalized_reason,
        )

    return students




async def create_student(
    db: AsyncSession,
    payload: StudentCreate,
    *,
    faculty_college: str,
    faculty_id: int | None = None,
    changed_by_faculty_id: int | None = None,
    changed_by_admin_id: int | None = None,
    require_hierarchy_faculty: bool = False,
) -> Student:
    faculty_college = (faculty_college or "").strip()

    if not faculty_college:
        raise ValueError("Faculty college is missing. Please set faculty.college.")

    name = payload.name.strip()
    usn = payload.usn.strip().upper()
    branch = payload.branch.strip()
    email = str(payload.email).strip().lower() if payload.email else None
    stype = _coerce_student_type(payload.student_type)
    required_points = _required_points_for_type(stype)

    dup_stmt = (
        select(Student)
        .where(
            Student.college == faculty_college,
            or_(
                Student.usn == usn,
                Student.email == email if email else False,
            ),
        )
        .order_by(Student.id.asc())
    )

    existing_students = (await db.execute(dup_stmt)).scalars().all()

    active_existing = next((s for s in existing_students if _is_active_student(s)), None)
    inactive_existing = next((s for s in existing_students if not _is_active_student(s)), None)

    if active_existing:
        if active_existing.usn == usn:
            raise ValueError(f"Duplicate USN in this college: {usn}")
        raise ValueError(f"Duplicate Email in this college: {email}")

    # --------------------------------------------------------
    # SERVER-DERIVED ACADEMIC BATCH
    # --------------------------------------------------------
    #
    # The client does not need to know an internal batch_id.
    # Admission year + passout year + college determine it.
    #
    # Example:
    #   GREEN CIRCUIT
    #   admitted_year = 2026
    #   passout_year  = 2030
    #       -> AcademicBatch 2026–2030
    #
    resolved_batch, default_current_year = (
        await resolve_student_academic_batch_from_years(
            db=db,
            college=faculty_college,
            admitted_year=payload.admitted_year,
            passout_year=payload.passout_year,
            requested_batch_id=payload.batch_id,
            requested_current_year=payload.current_year,
        )
    )

    academic_fields_set = {
        "department_id",
        "batch_id",
        "current_year",
        "assigned_faculty_id",
    }.intersection(payload.model_fields_set)

    if inactive_existing:
        previous_department_id = inactive_existing.department_id
        previous_batch_id = inactive_existing.batch_id
        previous_current_year = inactive_existing.current_year
        previous_assigned_faculty_id = (
            inactive_existing.assigned_faculty_id
        )

        target_department_id = (
            payload.department_id
            if "department_id" in payload.model_fields_set
            else previous_department_id
        )
        # Batch is always derived from admission/passout years.
        target_batch_id = int(
            resolved_batch.id
        )
        # Preserve an existing year only when the student remains
        # in the same resolved batch. Otherwise start from the
        # requested/default year.
        if payload.current_year is not None:
            target_current_year = int(
                payload.current_year
            )

        elif (
            previous_batch_id is not None
            and int(previous_batch_id)
            == int(resolved_batch.id)
            and previous_current_year is not None
        ):
            target_current_year = int(
                previous_current_year
            )

        else:
            target_current_year = int(
                default_current_year
            )
        target_assigned_faculty_id = (
            payload.assigned_faculty_id
            if "assigned_faculty_id" in payload.model_fields_set
            else previous_assigned_faculty_id
        )

        should_validate_assignment = bool(
            academic_fields_set
        )
    else:
        previous_department_id = None
        previous_batch_id = None
        previous_current_year = None
        previous_assigned_faculty_id = None

        target_department_id = payload.department_id
        target_batch_id = int(
            resolved_batch.id
        )

        target_current_year = (
            int(payload.current_year)
            if payload.current_year is not None
            else int(default_current_year)
        )
        target_assigned_faculty_id = (
            payload.assigned_faculty_id
        )

        should_validate_assignment = any(
            value is not None
            for value in (
                target_department_id,
                target_batch_id,
                target_current_year,
                target_assigned_faculty_id,
            )
        )

    if should_validate_assignment:
        await validate_student_academic_assignment(
            db=db,
            college=faculty_college,
            department_id=target_department_id,
            batch_id=target_batch_id,
            current_year=target_current_year,
            assigned_faculty_id=target_assigned_faculty_id,
            require_hierarchy_faculty=require_hierarchy_faculty,
        )

    # ✅ Permanent fix: restore deleted/deactivated student
    if inactive_existing:
        _reactivate_student(
            inactive_existing,
            college=faculty_college,
            name=name,
            email=email,
            usn=usn,
            branch=branch,
            student_type=stype,
            passout_year=payload.passout_year,
            admitted_year=payload.admitted_year,
            faculty_id=faculty_id,
        )

        inactive_existing.department_id = target_department_id
        inactive_existing.batch_id = target_batch_id
        inactive_existing.current_year = target_current_year
        inactive_existing.assigned_faculty_id = (
            target_assigned_faculty_id
        )

        record_student_assignment_history(
            db=db,
            student=inactive_existing,
            previous_department_id=previous_department_id,
            new_department_id=target_department_id,
            previous_batch_id=previous_batch_id,
            new_batch_id=target_batch_id,
            previous_year=previous_current_year,
            new_year=target_current_year,
            previous_faculty_id=previous_assigned_faculty_id,
            new_faculty_id=target_assigned_faculty_id,
            changed_by_faculty_id=changed_by_faculty_id,
            changed_by_admin_id=changed_by_admin_id,
            reason="Student reactivation academic assignment",
        )

        await db.commit()
        await db.refresh(inactive_existing)

        _trigger_student_welcome_email(inactive_existing.email, inactive_existing.name)

        return inactive_existing

    s = Student(
        college=faculty_college,
        name=name,
        usn=usn,
        branch=branch,
        email=email,
        student_type=stype,
        required_total_points=required_points,
        total_points_earned=0,
        passout_year=payload.passout_year,
        admitted_year=payload.admitted_year,
        department_id=target_department_id,
        batch_id=target_batch_id,
        current_year=target_current_year,
        assigned_faculty_id=target_assigned_faculty_id,
        created_by_faculty_id=faculty_id,
        is_active=True,
    )

    db.add(s)

    record_student_assignment_history(
        db=db,
        student=s,
        previous_department_id=None,
        new_department_id=target_department_id,
        previous_batch_id=None,
        new_batch_id=target_batch_id,
        previous_year=None,
        new_year=target_current_year,
        previous_faculty_id=None,
        new_faculty_id=target_assigned_faculty_id,
        changed_by_faculty_id=changed_by_faculty_id,
        changed_by_admin_id=changed_by_admin_id,
        reason="Initial student academic assignment",
    )

    await db.commit()
    await db.refresh(s)

    _trigger_student_welcome_email(s.email, s.name)

    return s


def convert_xlsx_to_csv_bytes(xlsx_bytes: bytes) -> bytes:
    """
    Convert the first worksheet of an .xlsx workbook into
    UTF-8 CSV bytes.

    The resulting bytes are intentionally passed through the
    existing create_students_from_csv() pipeline so CSV and
    Excel imports share the exact same validation, duplicate,
    hierarchy, batch, and ownership behavior.
    """
    if not xlsx_bytes:
        raise ValueError("Excel file is empty")

    try:
        workbook = load_workbook(
            filename=io.BytesIO(xlsx_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ValueError(
            "Unable to read Excel file. Please upload a valid .xlsx file."
        ) from exc

    try:
        worksheet = workbook.active

        if worksheet is None:
            raise ValueError(
                "Excel workbook does not contain a worksheet"
            )

        output = io.StringIO(newline="")
        writer = csv.writer(output)

        row_count = 0

        for excel_row in worksheet.iter_rows(values_only=True):
            normalized_row = []

            for value in excel_row:
                if value is None:
                    normalized_row.append("")
                    continue

                # Excel sometimes exposes whole-number cells as
                # floats. Keep years / IDs clean (2023, not 2023.0).
                if isinstance(value, float) and value.is_integer():
                    value = int(value)

                normalized_row.append(str(value).strip())

            # Preserve the header row and meaningful data rows while
            # ignoring completely empty trailing worksheet rows.
            if not any(normalized_row):
                continue

            writer.writerow(normalized_row)
            row_count += 1

        if row_count == 0:
            raise ValueError(
                "Excel file does not contain any rows"
            )

        return output.getvalue().encode("utf-8-sig")

    finally:
        workbook.close()


async def create_students_from_csv(
    db: AsyncSession,
    csv_bytes: bytes,
    skip_duplicates: bool = True,
    *,
    faculty_college: str,
    faculty_id: int | None = None,
    changed_by_faculty_id: int | None = None,
    changed_by_admin_id: int | None = None,
    allow_csv_college: bool = False,
    allow_csv_faculty_email: bool = False,
    hierarchy_role: str | None = None,
    hierarchy_department_id: int | None = None,
    hierarchy_faculty_id: int | None = None,
) -> Tuple[int, int, int, int, List[str]]:
    """
    Faculty upload CSV:
      name,email,usn,branch,student_type,admitted_year,passout_year

    Admin upload CSV:
      name,email,usn,branch,college,student_type,admitted_year,passout_year,faculty_email

    Optional academic headers for either flow:
      department_id,batch_id,current_year,assigned_faculty_id

    faculty_email controls creation ownership through
    created_by_faculty_id. assigned_faculty_id is a separate
    current academic assignment.

    Permanent duplicate behavior:
      - Active existing student => skip / duplicate
      - Inactive existing student => reactivate and update details
      - No existing student => insert new row
    """

    default_college = (faculty_college or "").strip()

    errors: List[str] = []
    inserted = 0
    skipped = 0
    invalid = 0
    welcome_targets: list[tuple[str, str]] = []

    try:
        text = csv_bytes.decode("utf-8-sig")
    except Exception:
        text = csv_bytes.decode("utf-8", errors="replace")

    reader = csv.DictReader(io.StringIO(text))

    if not reader.fieldnames:
        return (
            0,
            0,
            0,
            0,
            [
                "Import file has no headers. Required: "
                "name,email,usn,branch,student_type,admitted_year,passout_year"
            ],
        )

    field_map, headers = _normalize_csv_headers(reader.fieldnames)

    required = {"name", "email", "usn", "branch", "passout_year", "admitted_year"}

    if allow_csv_college:
        required.add("college")

    missing = required - headers

    if missing:
        return (0, 0, 0, 0, [f"Missing headers: {', '.join(sorted(missing))}"])

    rows = list(reader)
    total_rows = len(rows)

    # ✅ Permanent fix: preload full Student objects, not only email/usn.
    # We must know active vs inactive.
    existing_students = (await db.execute(select(Student))).scalars().all()

    active_usns: set[tuple[str, str]] = set()
    active_emails: set[tuple[str, str]] = set()

    inactive_by_usn: dict[tuple[str, str], Student] = {}
    inactive_by_email: dict[tuple[str, str], Student] = {}

    for s in existing_students:
        college_key = str(getattr(s, "college", "") or "").strip().lower()
        usn_key = str(getattr(s, "usn", "") or "").strip().upper()
        email_key = str(getattr(s, "email", "") or "").strip().lower()

        if not college_key:
            continue

        if _is_active_student(s):
            if usn_key:
                active_usns.add((college_key, usn_key))
            if email_key:
                active_emails.add((college_key, email_key))
        else:
            if usn_key:
                inactive_by_usn.setdefault((college_key, usn_key), s)
            if email_key:
                inactive_by_email.setdefault((college_key, email_key), s)

    faculty_by_email = {}

    if allow_csv_faculty_email and "faculty_email" in headers:
        faculty_rows = (await db.execute(select(Faculty))).scalars().all()
        faculty_by_email = {
            str(f.email or "").strip().lower(): f
            for f in faculty_rows
            if getattr(f, "email", None)
        }

    for idx, row in enumerate(rows, start=2):
        try:
            name = _clean(row.get(field_map["name"], ""))
            email = _clean(row.get(field_map["email"], "")).lower()
            usn = _clean(row.get(field_map["usn"], "")).upper()
            branch = _clean(row.get(field_map["branch"], ""))

            passout_year = int(_clean(row.get(field_map["passout_year"], "")))
            admitted_year = int(_clean(row.get(field_map["admitted_year"], "")))

            stype_raw = ""

            if "student_type" in headers:
                stype_raw = _clean(row.get(field_map["student_type"], "")).upper()

            if allow_csv_college and "college" in headers:
                row_college = _clean(row.get(field_map["college"], ""))
            else:
                row_college = default_college

            if not row_college:
                raise ValueError("college cannot be empty")

            row_faculty_id = faculty_id

            if allow_csv_faculty_email and "faculty_email" in headers:
                faculty_email = _clean(row.get(field_map["faculty_email"], "")).lower()

                if faculty_email:
                    faculty = faculty_by_email.get(faculty_email)

                    if not faculty:
                        raise ValueError(f"Faculty email not found: {faculty_email}")

                    row_faculty_id = faculty.id

                    if not row_college:
                        row_college = faculty.college

            if not name or not email or not usn or not branch:
                raise ValueError("name/email/usn/branch cannot be empty")

            college_key = row_college.strip().lower()
            usn_key = usn.strip().upper()
            email_key = email.strip().lower()

            active_dup = (
                (college_key, usn_key) in active_usns
                or (college_key, email_key) in active_emails
            )

            if active_dup:
                if skip_duplicates:
                    skipped += 1
                    continue

                raise ValueError("Duplicate USN/email in this college")

            stype = _parse_student_type(stype_raw)

            department_id = _parse_optional_csv_int(
                row=row,
                field_map=field_map,
                headers=headers,
                field_name="department_id",
            )
            batch_id = _parse_optional_csv_int(
                row=row,
                field_map=field_map,
                headers=headers,
                field_name="batch_id",
            )
            current_year = _parse_optional_csv_int(
                row=row,
                field_map=field_map,
                headers=headers,
                field_name="current_year",
                maximum=8,
            )
            assigned_faculty_id = _parse_optional_csv_int(
                row=row,
                field_map=field_map,
                headers=headers,
                field_name="assigned_faculty_id",
            )

            # CSV users do not need to know database batch IDs.
            # The active college batch is resolved from the
            # mandatory admitted_year + passout_year columns.
            resolved_batch, default_current_year = (
                await resolve_student_academic_batch_from_years(
                    db=db,
                    college=row_college,
                    admitted_year=admitted_year,
                    passout_year=passout_year,
                    requested_batch_id=batch_id,
                    requested_current_year=current_year,
                )
            )

            academic_headers = {
                "department_id",
                "batch_id",
                "current_year",
                "assigned_faculty_id",
            }.intersection(headers)

            # If the matching student is inactive, omitted CSV columns
            # normally preserve the existing academic assignment.
            inactive_student = (
                inactive_by_usn.get((college_key, usn_key))
                or inactive_by_email.get((college_key, email_key))
            )

            if inactive_student:
                previous_department_id = (
                    inactive_student.department_id
                )
                previous_batch_id = inactive_student.batch_id
                previous_current_year = (
                    inactive_student.current_year
                )
                previous_assigned_faculty_id = (
                    inactive_student.assigned_faculty_id
                )

                target_department_id = (
                    department_id
                    if "department_id" in headers
                    else previous_department_id
                )
                target_batch_id = int(
                    resolved_batch.id
                )
                if current_year is not None:
                    target_current_year = int(
                        current_year
                    )

                elif (
                    previous_batch_id is not None
                    and int(previous_batch_id)
                    == int(resolved_batch.id)
                    and previous_current_year is not None
                ):
                    target_current_year = int(
                        previous_current_year
                    )

                else:
                    target_current_year = int(
                        default_current_year
                    )
                target_assigned_faculty_id = (
                    assigned_faculty_id
                    if "assigned_faculty_id" in headers
                    else previous_assigned_faculty_id
                )
            else:
                previous_department_id = None
                previous_batch_id = None
                previous_current_year = None
                previous_assigned_faculty_id = None

                target_department_id = department_id
                target_batch_id = int(
                    resolved_batch.id
                )

                target_current_year = (
                    int(current_year)
                    if current_year is not None
                    else int(default_current_year)
                )
                target_assigned_faculty_id = assigned_faculty_id

            # -------------------------------------------------
            # Authenticated Faculty/HOD hierarchy enforcement.
            #
            # Admin bulk upload does not pass hierarchy_role and
            # therefore retains its existing full override behavior.
            # -------------------------------------------------
            if hierarchy_role == ROLE_FACULTY:
                if (
                    "department_id" in headers
                    and department_id is not None
                    and department_id != hierarchy_department_id
                ):
                    raise ValueError(
                        "Faculty/Mentor can bulk-create students "
                        "only inside their own department"
                    )

                if (
                    "assigned_faculty_id" in headers
                    and assigned_faculty_id is not None
                    and assigned_faculty_id != hierarchy_faculty_id
                ):
                    raise ValueError(
                        "Faculty/Mentor cannot assign a student "
                        "to another Faculty/Mentor"
                    )

                target_department_id = hierarchy_department_id

                # Preserve legacy Faculty behavior when no department
                # has yet been assigned. created_by_faculty_id remains
                # the legacy mentor provenance.
                target_assigned_faculty_id = (
                    hierarchy_faculty_id
                    if hierarchy_department_id is not None
                    else None
                )

            elif hierarchy_role == ROLE_HOD:
                if hierarchy_department_id is None:
                    raise ValueError(
                        "HOD department is not assigned"
                    )

                if (
                    "department_id" in headers
                    and department_id is not None
                    and department_id != hierarchy_department_id
                ):
                    raise ValueError(
                        "HOD can bulk-create students only "
                        "inside their own department"
                    )

                target_department_id = hierarchy_department_id

                # HOD assignment is explicit. If the CSV omits
                # assigned_faculty_id, do not preserve a historical
                # mentor that may be outside the HOD hierarchy.
                target_assigned_faculty_id = (
                    assigned_faculty_id
                    if "assigned_faculty_id" in headers
                    else None
                )

            should_validate_assignment = (
                hierarchy_role is not None
                or bool(academic_headers)
            ) and any(
                value is not None
                for value in (
                    target_department_id,
                    target_batch_id,
                    target_current_year,
                    target_assigned_faculty_id,
                )
            )

            if should_validate_assignment:
                await validate_student_academic_assignment(
                    db=db,
                    college=row_college,
                    department_id=target_department_id,
                    batch_id=target_batch_id,
                    current_year=target_current_year,
                    assigned_faculty_id=(
                        target_assigned_faculty_id
                    ),
                )

            # HOD may assign only an active Faculty/Mentor directly
            # beneath that authenticated HOD.
            if (
                hierarchy_role == ROLE_HOD
                and target_assigned_faculty_id is not None
            ):
                hierarchy_result = await db.execute(
                    select(Faculty).where(
                        Faculty.id == target_assigned_faculty_id,
                        Faculty.college == row_college,
                        Faculty.department_id
                        == hierarchy_department_id,
                        Faculty.parent_faculty_id
                        == hierarchy_faculty_id,
                        Faculty.role == ROLE_FACULTY,
                        Faculty.is_active.is_(True),
                    )
                )

                hierarchy_faculty = (
                    hierarchy_result.scalar_one_or_none()
                )

                if hierarchy_faculty is None:
                    raise ValueError(
                        "Selected Faculty/Mentor must be active, "
                        "belong to the HOD's department, and be "
                        "directly under this HOD"
                    )

            # ✅ Permanent fix:
            # If same USN/email exists but inactive, restore that
            # student instead of skipping.

            if inactive_student:
                _reactivate_student(
                    inactive_student,
                    college=row_college,
                    name=name,
                    email=email,
                    usn=usn,
                    branch=branch,
                    student_type=stype,
                    passout_year=passout_year,
                    admitted_year=admitted_year,
                    faculty_id=row_faculty_id,
                )

                inactive_student.department_id = (
                    target_department_id
                )
                inactive_student.batch_id = target_batch_id
                inactive_student.current_year = target_current_year
                inactive_student.assigned_faculty_id = (
                    target_assigned_faculty_id
                )

                record_student_assignment_history(
                    db=db,
                    student=inactive_student,
                    previous_department_id=(
                        previous_department_id
                    ),
                    new_department_id=target_department_id,
                    previous_batch_id=previous_batch_id,
                    new_batch_id=target_batch_id,
                    previous_year=previous_current_year,
                    new_year=target_current_year,
                    previous_faculty_id=(
                        previous_assigned_faculty_id
                    ),
                    new_faculty_id=(
                        target_assigned_faculty_id
                    ),
                    changed_by_faculty_id=(
                        changed_by_faculty_id
                    ),
                    changed_by_admin_id=changed_by_admin_id,
                    reason=(
                        "CSV student reactivation academic "
                        "assignment"
                    ),
                )

                inserted += 1

                active_usns.add((college_key, usn_key))
                active_emails.add((college_key, email_key))

                inactive_by_usn.pop((college_key, usn_key), None)
                inactive_by_email.pop((college_key, email_key), None)

                welcome_targets.append((email, name))
                continue

            required_points = _required_points_for_type(stype)

            s = Student(
                college=row_college,
                name=name,
                usn=usn,
                branch=branch,
                email=email,
                student_type=stype,
                required_total_points=required_points,
                total_points_earned=0,
                passout_year=passout_year,
                admitted_year=admitted_year,
                department_id=target_department_id,
                batch_id=target_batch_id,
                current_year=target_current_year,
                assigned_faculty_id=target_assigned_faculty_id,
                created_by_faculty_id=row_faculty_id,
                is_active=True,
            )

            db.add(s)

            record_student_assignment_history(
                db=db,
                student=s,
                previous_department_id=None,
                new_department_id=target_department_id,
                previous_batch_id=None,
                new_batch_id=target_batch_id,
                previous_year=None,
                new_year=target_current_year,
                previous_faculty_id=None,
                new_faculty_id=target_assigned_faculty_id,
                changed_by_faculty_id=changed_by_faculty_id,
                changed_by_admin_id=changed_by_admin_id,
                reason="Initial CSV student academic assignment",
            )

            inserted += 1

            active_usns.add((college_key, usn_key))
            active_emails.add((college_key, email_key))

            welcome_targets.append((email, name))

        except Exception as e:
            invalid += 1
            errors.append(f"Row {idx}: {str(e)}")

    await db.commit()

    for email, name in welcome_targets:
        _trigger_student_welcome_email(email, name)

    return (total_rows, inserted, skipped, invalid, errors)